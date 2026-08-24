"""Read-only structural inspection for the seven competition workbooks."""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import openpyxl
import xlrd


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_INPUT_DIR = ROOT / "C题数据"
DEFAULT_OUTPUT = ROOT / "supporting_materials" / "tables" / "audit" / "input_structure.json"


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def trim_row(values: list[Any]) -> list[Any]:
    last = -1
    for index, value in enumerate(values):
        if value not in (None, ""):
            last = index
    return [json_value(value) for value in values[: last + 1]]


def inspect_xlsx(path: Path, sample_rows: int) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        first_rows: list[list[Any]] = []
        last_nonempty: list[Any] = []
        nonempty_rows = 0
        formula_cells = 0
        for row_index, row in enumerate(worksheet.iter_rows(values_only=False), start=1):
            values = [cell.value for cell in row]
            trimmed = trim_row(values)
            if trimmed:
                nonempty_rows += 1
                last_nonempty = [row_index, *trimmed]
                if len(first_rows) < sample_rows:
                    first_rows.append([row_index, *trimmed])
            formula_cells += sum(
                isinstance(value, str) and value.startswith("=") for value in values
            )
        sheets.append(
            {
                "name": worksheet.title,
                "state": worksheet.sheet_state,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "nonempty_rows": nonempty_rows,
                "formula_cells": formula_cells,
                "first_nonempty_rows": first_rows,
                "last_nonempty_row": last_nonempty,
            }
        )
    workbook.close()
    return {"file": path.name, "format": "xlsx", "size_bytes": path.stat().st_size, "sheets": sheets}


def inspect_xls(path: Path, sample_rows: int) -> dict[str, Any]:
    workbook = xlrd.open_workbook(path, on_demand=True, formatting_info=False)
    sheets: list[dict[str, Any]] = []
    for name in workbook.sheet_names():
        worksheet = workbook.sheet_by_name(name)
        first_rows: list[list[Any]] = []
        last_nonempty: list[Any] = []
        nonempty_rows = 0
        formula_cells = None
        for row_index in range(worksheet.nrows):
            values = [worksheet.cell_value(row_index, col) for col in range(worksheet.ncols)]
            trimmed = trim_row(values)
            if trimmed:
                nonempty_rows += 1
                last_nonempty = [row_index + 1, *trimmed]
                if len(first_rows) < sample_rows:
                    first_rows.append([row_index + 1, *trimmed])
        sheets.append(
            {
                "name": name,
                "state": "visible",
                "max_row": worksheet.nrows,
                "max_column": worksheet.ncols,
                "nonempty_rows": nonempty_rows,
                "formula_cells": formula_cells,
                "first_nonempty_rows": first_rows,
                "last_nonempty_row": last_nonempty,
            }
        )
        workbook.unload_sheet(name)
    workbook.release_resources()
    return {"file": path.name, "format": "xls", "size_bytes": path.stat().st_size, "sheets": sheets}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--sample-rows", type=int, default=15)
    args = parser.parse_args()

    workbooks = sorted(
        [*args.input_dir.glob("*.xls"), *args.input_dir.glob("*.xlsx")],
        key=lambda item: item.name,
    )
    result = []
    for path in workbooks:
        print(f"Inspecting {path.name}", flush=True)
        if path.suffix.lower() == ".xls":
            result.append(inspect_xls(path, args.sample_rows))
        else:
            result.append(inspect_xlsx(path, args.sample_rows))

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
